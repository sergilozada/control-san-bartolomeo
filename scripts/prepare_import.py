"""Prepara los clientes del Excel sin inventar vencimientos ni pagos.

Uso: python scripts/prepare_import.py ruta/al/archivo.xlsx > import.json
El JSON contiene datos personales: no lo añadas al repositorio.
"""

import hashlib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def raw_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def amount(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = raw_text(value)
    if re.fullmatch(r"\d{1,3}(,\d{3})*(\.\d{1,2})?|\d+(\.\d{1,2})?", text):
        return float(text.replace(",", ""))
    return 0.0


def json_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def prepare(path):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    seen_lots = set()
    imported_at = datetime.now().astimezone().isoformat()
    for sheet in workbook:
        for row in sheet.iter_rows(min_row=2, values_only=False):
            values = [cell.value for cell in row]
            if not values or not raw_text(values[0]):
                continue
            values += [None] * max(0, 33 - len(values))
            name = raw_text(values[0])
            dni = raw_text(values[3])
            block = raw_text(values[7]).upper()
            lot = raw_text(values[8]).upper()
            lot_key = f"{block}|{lot}"
            if not block or not lot or lot_key in seen_lots:
                raise ValueError(f"Lote ausente o duplicado: {sheet.title}, fila {row[0].row}")
            seen_lots.add(lot_key)
            method_text = raw_text(values[11]).upper()
            method = "contado" if method_text == "CONTADO" else "cuotas"
            note = raw_text(values[14]) or raw_text(values[32])
            issues = [
                "El Excel no contiene fechas de vencimiento ni comprobantes de pago; cronograma pendiente de revisión.",
                "La hoja indica el mes, no la fecha exacta de registro o contrato.",
            ]
            if method_text not in {"CONTADO", "CREDITO", "CRÉDITO"}:
                issues.append("Forma de pago ausente o no reconocida; confirmar.")
            area = amount(values[9])
            total = amount(values[10])
            initial = amount(values[12])
            count = int(values[13]) if isinstance(values[13], (int, float)) and values[13] == int(values[13]) else 0
            if area == 0 or area > 1000:
                issues.append("Metraje ambiguo en la fuente; verificar unidad y cifra antes de emitir documentos.")
            if total <= 0:
                issues.append("Precio total ausente o con caracteres ambiguos; confirmar.")
            if method == "cuotas" and (count < 1 or (values[12] is not None and initial == 0)):
                issues.append("Inicial o número de cuotas incompleto/ambiguo; confirmar.")
            if not re.fullmatch(r"\d{8}", dni):
                issues.append("DNI requiere revisión; se conservó literalmente del Excel.")
            if note:
                issues.append("Hay una observación contractual o de pago que requiere revisión humana.")
            original = {
                "hoja": sheet.title.strip(),
                "nombre": json_value(values[0]),
                "dni": json_value(values[3]),
                "direccion": json_value(values[4]),
                "telefono": json_value(values[5]),
                "email": json_value(values[6]),
                "manzana": json_value(values[7]),
                "lote": json_value(values[8]),
                "metraje": json_value(values[9]),
                "precio": json_value(values[10]),
                "formaPago": json_value(values[11]),
                "inicial": json_value(values[12]),
                "numeroCuotas": json_value(values[13]),
                "observacion": note or None,
            }
            doc_id = "san-" + hashlib.sha256(lot_key.encode()).hexdigest()[:20]
            record = {
                "id": doc_id,
                "nombre1": name,
                "dni1": dni,
                "titulares": [{"nombre": name, "dni": dni}],
                "celular1": raw_text(values[5]),
                "email1": raw_text(values[6]),
                "manzana": block,
                "lote": lot,
                "metraje": area,
                "montoTotal": total,
                "formaPago": method,
                "inicial": initial if method == "cuotas" else 0,
                "numeroCuotas": count if method == "cuotas" else 1,
                "fechaRegistro": date.today().isoformat(),
                "cuotas": [],
                "importReview": {
                    "status": "pending",
                    "sourceFile": path.name,
                    "sourceHash": digest,
                    "sourceRow": row[0].row,
                    "issues": issues,
                    "original": original,
                    "dateMapping": "sheet-month-only; import-date-is-not-contract-date",
                },
            }
            if note:
                record["observationEntries"] = [{
                    "id": "excel-" + hashlib.sha256((doc_id + note).encode()).hexdigest()[:16],
                    "text": note,
                    "author": "Importación Excel · " + sheet.title.strip() + " · fila " + str(row[0].row),
                    "at": imported_at,
                }]
            records.append(record)
    return {"sourceHash": digest, "count": len(records), "clients": records}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Indica la ruta del Excel.")
    json.dump(prepare(Path(sys.argv[1])), sys.stdout, ensure_ascii=False, separators=(",", ":"))
